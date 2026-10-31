# -*- coding: utf-8 -*-
"""
课题03 - 图表生成器（第二轮 · RSSI 口径）

产出：
  charts/chart1_RSSI对比.png        各点位 RSSI 均值（-80 dBm 橙色虚线，跌破标红）
  charts/chart2_延迟对比.png        各点位下行延迟均值（无网络点位标灰）
  charts/chart3_信号与速率散点.png   RSSI vs 下载速率
  charts/chart4_楼层对比.png        2/5/8 楼电梯口对照
  charts/课题03_图表.xlsx           Excel 原生图表

只用 openpyxl + Pillow，不需要 matplotlib。
"""

import csv
import io
import math
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import StrRef

from PIL import Image, ImageDraw, ImageFont

# ---------- 路径定位：兼容工作区布局与仓库布局 ----------
_here = Path(__file__).resolve().parent
_cands = [_here.parent]
if _here.parent.parent.name == "01-Field-Discovery":
    _cands.append(_here.parent.parent)
BASE = None
for _c in _cands:
    if (_c / "raw_data").is_dir() and (_c / "photos").is_dir():
        BASE = _c
        break
if BASE is None:
    BASE = _here.parent
RAW = BASE / "raw_data"
if not (RAW / "measurements.csv").exists():
    for _sub in RAW.iterdir():
        if _sub.is_dir() and (_sub / "measurements.csv").exists():
            RAW = _sub
            break
CHARTS = BASE / "charts"
CHARTS.mkdir(exist_ok=True)
OUT = CHARTS / "课题03_图表.xlsx"

BLIND_DBM = -80
C_BLUE, C_RED, C_ORANGE = "#2E75B6", "#C00000", "#FF6600"

W, H = 1600, 980
PAD_L, PAD_R, PAD_T, PAD_B = 150, 60, 110, 200

FONTS = [("C:/Windows/Fonts/msyh.ttc", 0), ("C:/Windows/Fonts/simhei.ttf", 0),
         ("C:/Windows/Fonts/simsun.ttc", 0)]


def font(size):
    for p, i in FONTS:
        if Path(p).exists():
            try:
                return ImageFont.truetype(p, size, index=i)
            except Exception:
                continue
    return ImageFont.load_default()


def num(text):
    s = (text or "").strip()
    if not s or s in ("无网络", "无服务", "无读数", "未连接", "超时", "-", "/"):
        return None
    try:
        return float(re.findall(r"-?\d+\.?\d*", s)[0])
    except IndexError:
        return None


def load(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "gbk", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SystemExit(f"无法识别 {path} 的编码")
    return list(csv.DictReader(io.StringIO(text, newline="")))


def avg(vals):
    v = [x for x in vals if x is not None]
    return sum(v) / len(v) if v else None


def axes(draw, img, labels, y_min, y_max, ticks, y_title, x_title, title):
    x0, y0 = PAD_L, PAD_T
    x1, y1 = W - PAD_R, H - PAD_B
    f_tick, f_title, f_head = font(24), font(26), font(38)
    draw.text(((x0 + x1) // 2, 45), title, font=f_head, fill="#1F3864", anchor="mm")
    span = y_max - y_min
    for v in ticks:
        y = y1 - (v - y_min) / span * (y1 - y0)
        draw.line([(x0, y), (x1, y)], fill="#D9D9D9", width=2)
        draw.text((x0 - 15, y), f"{v:g}", font=f_tick, fill="#404040", anchor="rm")
    draw.line([(x0, y0), (x0, y1)], fill="#404040", width=3)
    draw.line([(x0, y1), (x1, y1)], fill="#404040", width=3)
    txt = Image.new("RGBA", (420, 60), (0, 0, 0, 0))
    d = ImageDraw.Draw(txt)
    d.text((210, 30), y_title, font=f_title, fill="#1F3864", anchor="mm")
    txt = txt.rotate(90, expand=True)
    img.paste(txt, (20, (y0 + y1) // 2 - 210), txt)
    draw.text(((x0 + x1) // 2, H - 38), x_title, font=f_title, fill="#1F3864", anchor="mm")
    n = len(labels)
    step = (x1 - x0) / max(n, 1)
    for i, (pid, name) in enumerate(labels):
        cx = x0 + step * (i + 0.5)
        draw.text((cx, y1 + 28), pid, font=f_title, fill="#1F3864", anchor="mm")
        for k, ch in enumerate([name[j:j + 6] for j in range(0, len(name), 6)]):
            draw.text((cx, y1 + 58 + k * 26), ch, font=f_tick, fill="#404040", anchor="mm")
    return x0, y0, x1, y1, step


def nice_step(vmax, target=5):
    raw = max(vmax, 1e-6) / target
    mag = 10 ** math.floor(math.log10(raw))
    for m in (1, 2, 2.5, 5, 10):
        if m * mag >= raw:
            return m * mag
    return 10 * mag


def chart_rssi(points, rows, out):
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    labels, vals = [], []
    for p in points:
        pid = p["点位编号"]
        labels.append((pid, p["点位名称"]))
        vals.append(avg([num(r["信号强度RSSI(dBm)"]) for r in rows if r["点位编号"] == pid]))
    x0, y0, x1, y1, step = axes(d, img, labels, -100, 0, [0, -20, -40, -60, -80, -100],
                                "RSSI (dBm)", "点位",
                                "课题03 · 各点位信号强度 RSSI 均值（红柱 = 盲区）")
    bw = step * 0.55
    for i, v in enumerate(vals):
        cx = x0 + step * (i + 0.5)
        if v is None:
            d.text((cx, y1 - 34), "无读数", font=font(21), fill="#7F7F7F", anchor="mm")
            continue
        top = y1 - (v + 100) / 100 * (y1 - y0)
        d.rectangle([cx - bw / 2, top, cx + bw / 2, y1],
                    fill=C_RED if v < BLIND_DBM else C_BLUE)
        d.text((cx, top - 18), f"{v:.1f}", font=font(22), fill="#1F1F1F", anchor="mm")
    yb = y1 - (BLIND_DBM + 100) / 100 * (y1 - y0)
    for xs in range(x0, x1, 26):
        d.line([(xs, yb), (min(xs + 14, x1), yb)], fill=C_ORANGE, width=4)
    # 阈值标签放在最右侧空白区（P8/P10 无柱），避免压住柱子
    d.text((x1 - 14, yb - 16), f"盲区阈值 {BLIND_DBM} dBm", font=font(22),
           fill=C_ORANGE, anchor="rs")
    img.save(out)


def chart_latency(points, rows, out):
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    labels, vals = [], []
    for p in points:
        pid = p["点位编号"]
        labels.append((pid, p["点位名称"]))
        vals.append(avg([num(r["下行延迟ms"]) for r in rows if r["点位编号"] == pid]))
    vmax = max([v for v in vals if v] + [100]) * 1.2
    s = nice_step(vmax)
    x0, y0, x1, y1, step = axes(d, img, labels, 0, s * 5, [s * i for i in range(6)],
                                "下行延迟 (ms)", "点位",
                                "课题03 · 各点位下行延迟均值（红 = 超过 500ms）")
    bw = step * 0.55
    for i, v in enumerate(vals):
        cx = x0 + step * (i + 0.5)
        if v is None:
            d.text((cx, y1 - 34), "无网络", font=font(21), fill="#7F7F7F", anchor="mm")
            continue
        top = y1 - v / (s * 5) * (y1 - y0)
        d.rectangle([cx - bw / 2, top, cx + bw / 2, y1],
                    fill=C_RED if v > 500 else "#ED7D31")
        d.text((cx, top - 18), f"{v:.0f}", font=font(22), fill="#1F1F1F", anchor="mm")
    img.save(out)


def chart_scatter(points, rows, out):
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    pts = []
    for p in points:
        for r in rows:
            if r["点位编号"] != p["点位编号"]:
                continue
            x, y = num(r["信号强度RSSI(dBm)"]), num(r["下载Mbps"])
            if x is not None and y is not None:
                pts.append((x, y))
    x_min, x_max = -90, -35
    s = nice_step(max([q[1] for q in pts] + [50]) * 1.15)
    y_max = s * 5
    f_tick, f_title, f_head = font(24), font(26), font(36)
    x0, y0, x1, y1 = PAD_L, PAD_T, W - PAD_R, H - PAD_B
    d.text(((x0 + x1) // 2, 45), "课题03 · 信号强度与下载速率的对应关系",
           font=f_head, fill="#1F3864", anchor="mm")
    for i in range(6):
        y = y1 - (s * i) / y_max * (y1 - y0)
        d.line([(x0, y), (x1, y)], fill="#D9D9D9", width=2)
        d.text((x0 - 15, y), f"{s*i:g}", font=f_tick, fill="#404040", anchor="rm")
    for v in (-90, -80, -70, -60, -50, -40):
        x = x0 + (v - x_min) / (x_max - x_min) * (x1 - x0)
        d.line([(x, y0), (x, y1)], fill="#EFEFEF", width=2)
        d.text((x, y1 + 28), str(v), font=f_tick, fill="#404040", anchor="mm")
    d.line([(x0, y0), (x0, y1)], fill="#404040", width=3)
    d.line([(x0, y1), (x1, y1)], fill="#404040", width=3)
    xb = x0 + (BLIND_DBM - x_min) / (x_max - x_min) * (x1 - x0)
    for ys in range(y0, y1, 26):
        d.line([(xb, ys), (xb, min(ys + 14, y1))], fill=C_ORANGE, width=4)
    d.text((xb + 10, y0 + 14), f"盲区线 {BLIND_DBM} dBm", font=font(22), fill=C_ORANGE)
    for x, y in pts:
        px = x0 + (x - x_min) / (x_max - x_min) * (x1 - x0)
        py = y1 - y / y_max * (y1 - y0)
        d.ellipse([px - 11, py - 11, px + 11, py + 11], fill=C_BLUE, outline="white", width=2)
    txt = Image.new("RGBA", (440, 60), (0, 0, 0, 0))
    dd = ImageDraw.Draw(txt)
    dd.text((220, 30), "下载速率 (Mbps)", font=f_title, fill="#1F3864", anchor="mm")
    rot = txt.rotate(90, expand=True)
    img.paste(rot, (20, (y0 + y1) // 2 - 220), rot)
    d.text(((x0 + x1) // 2, H - 38), "信号强度 RSSI (dBm)", font=f_title, fill="#1F3864", anchor="mm")
    img.save(out)


def chart_floor(rows, out):
    trio = [("P7", "2楼电梯口"), ("P5", "5楼电梯口"), ("P6", "8楼电梯口")]
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    labels, vals = [], []
    for pid, name in trio:
        labels.append((pid, name))
        vals.append(avg([num(r["信号强度RSSI(dBm)"]) for r in rows if r["点位编号"] == pid]))
    x0, y0, x1, y1, step = axes(d, img, labels, -100, 0, [0, -20, -40, -60, -80, -100],
                                "RSSI (dBm)", "楼层对照点",
                                "课题03 · 不同楼层电梯口信号强度对照")
    bw = step * 0.42
    for i, v in enumerate(vals):
        if v is None:
            continue
        cx = x0 + step * (i + 0.5)
        top = y1 - (v + 100) / 100 * (y1 - y0)
        d.rectangle([cx - bw / 2, top, cx + bw / 2, y1], fill=C_RED if v < BLIND_DBM else C_BLUE)
        d.text((cx, top - 24), f"{v:.1f}", font=font(26), fill="#1F1F1F", anchor="mm")
    yb = y1 - (BLIND_DBM + 100) / 100 * (y1 - y0)
    for xs in range(x0, x1, 26):
        d.line([(xs, yb), (min(xs + 14, x1), yb)], fill=C_ORANGE, width=4)
    d.text((x0 + 12, yb - 18), f"盲区阈值 {BLIND_DBM} dBm", font=font(22), fill=C_ORANGE)
    img.save(out)


def main():
    rows = load(RAW / "measurements.csv")
    points = load(RAW / "pose_points.csv")

    chart_rssi(points, rows, CHARTS / "chart1_RSSI对比.png")
    chart_latency(points, rows, CHARTS / "chart2_延迟对比.png")
    chart_scatter(points, rows, CHARTS / "chart3_信号与速率散点.png")
    chart_floor(rows, CHARTS / "chart4_楼层对比.png")

    wb = Workbook()
    ws = wb.active
    ws.title = "RSSI"
    ws.append(["点位", "RSSI均值(dBm)"])
    for p in points:
        pid = p["点位编号"]
        v = avg([num(r["信号强度RSSI(dBm)"]) for r in rows if r["点位编号"] == pid])
        ws.append([p["点位名称"], v if v is not None else -100])
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 16
    ch = BarChart()
    ch.type = "col"
    ch.title = "课题03 · 各点位 RSSI 均值"
    ch.y_axis.title = "RSSI (dBm)"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    try:
        ch.series[0].cat.strRef = StrRef(f"'RSSI'!$A$2:$A${ws.max_row}")
    except Exception:
        pass
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.scaling.min = -100
    ch.y_axis.scaling.max = 0
    ch.legend = None
    ch.height, ch.width = 10, 24
    ws.add_chart(ch, "D2")
    wb.save(OUT)

    n = len([p for p in points if avg([num(r["信号强度RSSI(dBm)"]) for r in rows
                                       if r["点位编号"] == p["点位编号"]]) is not None])
    print(f"[OK] charts 4 png + xlsx | points={len(points)} records={len(rows)} rssi_ok={n}")


if __name__ == "__main__":
    main()